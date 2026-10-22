#!/usr/bin/env python3
"""Build decade.json: the 2010s Decade-End lists from Will's "CS Era points - 2010s Decade" workbook.

    python3 tools/build_decade.py "CS Era points - 2010s Decade.xlsx"

Run from the repo folder. Country songs, artists and albums keep the order of Will's
"Top Songs", "Top Artists" and "Top Albums" tabs. The Pop tabs aren't ranked, so Pop
songs and artists are ordered by their decade chart points, the way Will ranks Pop.
Points only set the order; they aren't written out, since the site never shows them.
"""
import json
import re
import sys
from collections import defaultdict

import openpyxl

from add_week import norm
from build_yearend import match

DECADE, TOP_SONGS, TOP_OTHER = '2010s', 100, 25
RENAMED = {'Lady Antebellum': 'Lady A'}  # the name the rest of the site uses


def rows(wb, tab):
    it = wb[tab].iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else '' for h in next(it)]
    return [dict(zip(head, r)) for r in it]


def clean(v):
    return re.sub(r'\s+', ' ', str(v or '')).strip()


def main():
    wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
    data = json.load(open('data.json', encoding='utf-8'))
    runs = json.load(open('runs.json', encoding='utf-8'))
    albums = json.load(open('albums.json', encoding='utf-8'))
    artist_id = {norm(a['name']): a['id'] for a in data['artists']}
    out, unmatched = {}, []

    def song_rows(key, listed):
        songs = [s for s in data['songs'] if s['chart'] == key]
        by_name = defaultdict(list)
        for s in songs:
            by_name[norm(s['name'])].append(s)
        res = []
        for title, artist, year in listed:
            s = match(by_name, title, artist, year)
            run = [p for p in runs[key]['runs'].get(s['id'], [None])[1:] if p] if s else []
            if not s:
                unmatched.append(f'{key}: {title} - {artist}')
            res.append({'rank': len(res) + 1, 'songId': s['id'] if s else None,
                        'name': s['name'] if s else title, 'artistName': s['artistName'] if s else artist,
                        'peak': min(run) if run else None, 'wks1': run.count(1), 'wc': len(run)})
            if len(res) == TOP_SONGS:
                break
        return res

    def artist_rows(names):
        # The Pop tab lists some acts twice and has pairings ("Maroon 5; Wiz Khalifa") of its own
        names = list(dict.fromkeys(RENAMED.get(n, n) for n in names if ';' not in n))[:TOP_OTHER]
        return [{'rank': i + 1, 'name': n, 'artistId': artist_id.get(norm(n))} for i, n in enumerate(names)]

    # Country: Will's own order
    top = [r for r in rows(wb, 'Top Songs') if r.get('Song')]
    country = song_rows('willboard', [(clean(r['Song']), clean(r['Artist']), getattr(r['Debut Date'], 'year', 2015)) for r in top])
    album_artist = {}
    for r in top:
        album_artist.setdefault(norm(clean(r['Album'])), clean(r['Artist']))
    al_rows = []
    for r in rows(wb, 'Top Albums'):
        name = clean(r.get('Album'))
        if not name or not r.get('Total Singles'):  # a stray duplicate row with no singles
            continue
        artist = album_artist.get(norm(name), '')
        artist = RENAMED.get(artist, artist)
        hit = next((a for a in albums if norm(a['name']) == norm(name)
                    and (not artist or norm(a['artist']).startswith(norm(artist)[:6]))), None)
        al_rows.append({'rank': len(al_rows) + 1, 'name': hit['name'] if hit else name,
                        'artist': hit['artist'] if hit else artist})
        if len(al_rows) == TOP_OTHER:
            break
    out['willboard'] = {'songs': country,
                        'artists': artist_rows([clean(r['Artist']) for r in rows(wb, 'Top Artists') if r.get('Artist')]),
                        'albums': al_rows}

    # Pop: ordered by decade points
    it = wb['Pop'].iter_rows(values_only=True)
    next(it)
    head = [str(h).strip() if h is not None else '' for h in next(it)]
    pop = [dict(zip(head, r)) for r in it]
    pop = sorted((r for r in pop if r.get('Song') and isinstance(r.get('Total'), (int, float))), key=lambda r: -r['Total'])
    pop_songs = song_rows('pop', [(clean(r['Song']), clean(r['Artist']).replace(';', ','), getattr(r['Debut Date'], 'year', 2015)) for r in pop])
    pop_artists = sorted((r for r in rows(wb, 'Pop Artist Records') if r.get('Artist') and isinstance(r.get('Points'), (int, float))),
                         key=lambda r: -r['Points'])
    out['pop'] = {'songs': pop_songs, 'artists': artist_rows([clean(r['Artist']) for r in pop_artists])}

    json.dump({DECADE: out}, open('decade.json', 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    for k, v in out.items():
        print(k, {n: len(x) for n, x in v.items()}, 'top:', [s['name'] for s in v['songs'][:3]])
    print('Unmatched:', *unmatched, sep='\n  ')


if __name__ == '__main__':
    main()
