#!/usr/bin/env python3
"""Build week_plays.json: each charting song's plays in every chart week we have them for.

    python3 tools/build_week_plays.py <Plays Pages folder> <country_plays.csv> <pop_plays.csv>

Run from the repo folder. The Plays Pages folder holds Will's "Plays Page YYYY" workbooks
saved as YYYY.xlsx (2015-2019). Sources, per Will (2026-10-26):
  Hot 20  2015-2019  Plays Pages: a block per week of song, start, end, plays, position
          2020 on    CS Era Points plays archive: running total and position per week
  Pop 10  2018-2019  Plays Pages "Pop Plays" tab: plays per song per week (Last.fm)
          2020 on    CS Era Points pop plays archive: plays per song per week
Hot 20 weeks store [plays that week, running total]; the running total drives the
certification shown on old chart pages. Pop 10 weeks store plays that week only (Pop 10
certifications come from chart points). Plays Page blocks are matched to their chart week by
the songs and positions in them, since Will's chart numbers there don't line up with dates.
"""
import csv
import datetime
import difflib
import json
import os
import sys
from collections import defaultdict

import openpyxl

from build_year_tops import is_num, keys
from build_data import WEEK_LABEL

MONTHS = 'Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec'.split()


def week_entries(runs, key):
    """{week: {position: songId}} from runs.json."""
    r = runs[key]
    out = defaultdict(dict)
    for sid, run in r['runs'].items():
        for i, p in enumerate(run[1:]):
            if p:
                out[r['weeks'][run[0] + i]][p] = sid
    return out


def matcher(songs):
    """Match a title to one of a week's songs by name, loosest form last."""
    def find(title, sids):
        by = {}
        for sid in sids:
            for k in keys(songs[sid]['name']):
                by.setdefault(k, sid)
        for k in keys(title):
            if k in by:
                return by[k]
        for k in keys(title):
            hit = difflib.get_close_matches(k, list(by), n=1, cutoff=0.8)
            if hit:
                return by[hit[0]]
        return None
    return find


def plays_page_blocks(path, tab):
    """Every week block on a Hot 20 Plays Page: [(song, start, end, plays, position)]."""
    rows = list(openpyxl.load_workbook(path, data_only=True)[tab].iter_rows(values_only=True))
    blocks = defaultdict(list)  # (band start row, column) -> records
    band = 0
    for i, r in enumerate(rows):
        r = list(r)
        recs = [(c, (r[c], int(r[c + 1]), int(r[c + 2]), int(r[c + 3]), int(r[c + 4])))
                for c in range(len(r) - 4) if isinstance(r[c], str) and all(is_num(v) for v in r[c + 1:c + 5])]
        if not recs:
            band = i  # a header or blank row starts a new band of weeks
        for c, rec in recs:
            blocks[(band, c)].append(rec)
    return [b for b in blocks.values() if len(b) >= 10]


def main():
    pp, country_csv, pop_csv = sys.argv[1:4]
    data = json.load(open('data.json', encoding='utf-8'))
    runs = json.load(open('runs.json', encoding='utf-8'))
    songs = {s['id']: s for s in data['songs']}
    find = matcher(songs)
    out = {'willboard': defaultdict(dict), 'pop': defaultdict(dict)}
    missed = defaultdict(int)

    # Hot 20, 2015-2019: match each block to the week whose chart it reproduces
    wb = week_entries(runs, 'willboard')
    for y in range(2015, 2020):
        path = os.path.join(pp, f'{y}.xlsx')
        if not os.path.exists(path):
            continue
        tab = next(t for t in openpyxl.load_workbook(path, read_only=True).sheetnames if t.startswith('Country') and 'Total' not in t and 'Points' not in t)
        for block in plays_page_blocks(path, tab):
            best, score = None, 0
            for w, ents in wb.items():
                if not (f'{y - 1}-12' <= w <= f'{y}-12-31'):
                    continue
                s = sum(1 for song, *_, pos in block if ents.get(pos) and find(song, [ents[pos]]))
                if s > score:
                    best, score = w, s
            if not best or score < len(block) // 2:
                missed['hot20 block'] += 1
                continue
            ents = wb[best]
            for song, start, end, plays, pos in block:
                sid = find(song, list(ents.values()))
                if sid:
                    out['willboard'][best][sid] = [plays, end]
                else:
                    missed['hot20 song'] += 1

    # Hot 20, 2020 on: the archive's running totals, a (total, position) column pair per week.
    # The pairs drift against the calendar (skipped and repeated weeks), so each position column
    # is matched to the chart week whose songs sit at those positions. Weeks from data.json's
    # first week on already carry their plays there.
    rows = list(csv.reader(open(country_csv, encoding='utf-8')))
    recent = data['willboard'][0]['week']
    span = [w for w in wb if '2019-12-26' <= w < recent and w not in out['willboard']]
    best = {}  # week -> (score, position column)
    for c in range(2, max(map(len, rows))):
        tally = defaultdict(int)
        for r in rows[1:]:
            if c < len(r) and r[c].strip().isdigit():
                pos = int(r[c])
                for w in span:
                    if wb[w].get(pos) and find(r[0], [wb[w][pos]]):
                        tally[w] += 1
        if tally:
            w = max(tally, key=tally.get)
            if tally[w] >= 12 and tally[w] > best.get(w, (0,))[0]:
                best[w] = (tally[w], c)
    cols = sorted((c, w) for w, (_, c) in best.items())
    for k, (c, w) in enumerate(cols):
        before = cols[k - 1][0] - 1 if k and cols[k - 1][1] < w else None
        for r in rows[1:]:
            if c >= len(r) or not r[c].strip().isdigit() or not r[c - 1].strip().isdigit():
                continue
            pos, total = int(r[c]), int(r[c - 1])
            sid = wb[w].get(pos) if wb[w].get(pos) and find(r[0], [wb[w][pos]]) else find(r[0], list(wb[w].values()))
            if not sid:
                continue
            was = r[before].strip() if before is not None and before < len(r) else ''
            out['willboard'][w][sid] = [total - int(was) if was.isdigit() else None, total]
    missed['hot20 archive weeks unmatched'] = len(set(span) - set(best))

    # Pop 10, 2018-2019: Plays Pages "Pop Plays" tabs
    pop = week_entries(runs, 'pop')
    def add_pop(w, title, plays):
        if w not in pop or not plays:
            return
        sid = find(title, list(pop[w].values()))
        if sid:
            out['pop'][w][sid] = int(plays)
    for y in (2018, 2019):
        path = os.path.join(pp, f'{y}.xlsx')
        if not os.path.exists(path):
            continue
        rows = list(openpyxl.load_workbook(path, data_only=True)['Pop Plays'].iter_rows(values_only=True))
        cols = {j: v.date().isoformat() for j, v in enumerate(rows[0]) if isinstance(v, datetime.datetime)}
        for r in rows[1:]:
            if not r[0]:
                continue
            for j, w in cols.items():
                if j < len(r) and is_num(r[j]):
                    add_pop(w, r[0], r[j])

    # Pop 10, 2020 on: the archive's weekly plays (dates are "2-Jan" labels, year rolls over in order)
    rows = list(csv.reader(open(pop_csv, encoding='utf-8')))
    # 2025's columns reuse 2024's labels (29-Feb, no 22-Aug), so a year whose labels aren't
    # Thursdays is read in order against that year's chart weeks instead
    year, prev, by_year = 2020, 0, defaultdict(list)
    for j, h in enumerate(rows[0]):
        m = WEEK_LABEL.fullmatch(h.strip()) if j else None
        if not m:
            continue
        mon = MONTHS.index(m.group(2)) + 1
        if mon < prev:
            year += 1
        prev = mon
        by_year[year].append((j, mon, int(m.group(1))))
    cols = {}
    for year, labels in by_year.items():
        chart = [w for w in runs['pop']['weeks'] if w.startswith(str(year))]
        try:
            dates = [datetime.date(year, mon, day).isoformat() for _, mon, day in labels]
        except ValueError:
            dates = []
        if dates and all(datetime.date.fromisoformat(d).weekday() == 3 for d in dates):
            cols.update({j: d for (j, *_), d in zip(labels, dates) if d in chart})
        elif len(labels) == len(chart):
            cols.update({j: d for (j, *_), d in zip(labels, chart)})
        else:
            missed[f'pop {year} columns'] += 1
    for j, w in cols.items():
        for r in rows[1:]:
            if j < len(r) and r[j].strip().isdigit():
                add_pop(w, r[0], r[j])

    res = {k: {w: v[w] for w in sorted(v)} for k, v in out.items()}
    json.dump(res, open('week_plays.json', 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    for k, v in res.items():
        print(k, len(v), 'weeks', min(v), '..', max(v), sum(len(x) for x in v.values()), 'song-weeks')
    print('missed', dict(missed))


if __name__ == '__main__':
    main()
